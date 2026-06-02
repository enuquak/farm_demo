# -*- coding: utf-8 -*-
"""Excel 配置表读取器。

读取符合规范的 Excel 文件，解析字段名、类型、注释和数据行。

Excel 格式约定：
- 行 1：字段名（snake_case）
- 行 2：字段类型（int/float/str/bool）
- 行 3：字段描述（注释）
- 行 4 起：数据
"""
import os
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import load_workbook

# 支持的字段类型及对应的 Python 类型转换函数
TYPE_CONVERTERS = {
    "int": int,
    "float": float,
    "str": str,
    "bool": bool,
}

SUPPORTED_TYPES = set(TYPE_CONVERTERS.keys())


def read_excel(file_path: str) -> Dict[str, Any]:
    """读取 Excel 配置文件，返回结构化数据。

    Args:
        file_path: Excel 文件路径。

    Returns:
        dict，包含以下键：
        - table_name: str，表名（文件名去扩展名）
        - fields: list[str]，字段名列表
        - types: list[str]，字段类型列表
        - comments: list[str]，字段描述列表
        - rows: list[list]，数据行列表（已做类型转换）

    Raises:
        FileNotFoundError: 文件不存在。
        ValueError: 字段类型不支持或 Excel 格式错误。
    """
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"Excel 文件不存在: {file_path}")

    wb = load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active

    # 读取行 1：字段名
    fields = [cell.value for cell in ws[1]]
    if not fields or all(f is None for f in fields):
        raise ValueError(f"{file_path}: 行 1（字段名）为空")

    # 读取行 2：字段类型
    types = [cell.value for cell in ws[2]]
    if len(types) != len(fields):
        raise ValueError(f"{file_path}: 行 2（字段类型）列数与行 1 不匹配")

    # 验证类型
    for i, t in enumerate(types):
        if t is None:
            raise ValueError(f"{file_path}: 行 2, 第 {i+1} 列类型为空")
        if str(t).lower() not in SUPPORTED_TYPES:
            raise ValueError(
                f"{file_path}: 不支持的字段类型 '{t}'，"
                f"支持的类型: {', '.join(sorted(SUPPORTED_TYPES))}"
            )

    # 读取行 3：注释
    comments = [cell.value for cell in ws[3]]

    # 读取行 4 起：数据行
    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        # 跳过全空行
        if all(cell is None for cell in row):
            continue
        converted = _convert_row(row, types, file_path, row_idx)
        rows.append(converted)

    wb.close()

    table_name = Path(file_path).stem
    return {
        "table_name": table_name,
        "fields": fields,
        "types": [str(t).lower() for t in types],
        "comments": comments,
        "rows": rows,
    }


def _convert_row(
    row: tuple, types: List[str], file_path: str, row_idx: int
) -> List[Any]:
    """将一行原始数据按类型转换。

    Args:
        row: 原始数据行（tuple）。
        types: 字段类型列表。
        file_path: 文件路径（用于错误信息）。
        row_idx: 行号（用于错误信息）。

    Returns:
        转换后的数据列表。

    Raises:
        ValueError: 类型转换失败。
    """
    result = []
    for col_idx, (value, type_name) in enumerate(zip(row, types)):
        if value is None:
            result.append(None)
            continue
        converter = TYPE_CONVERTERS[type_name.lower()]
        try:
            result.append(converter(value))
        except (ValueError, TypeError) as e:
            raise ValueError(
                f"{file_path}: 第 {row_idx} 行, 第 {col_idx+1} 列: "
                f"无法将 '{value}' 转换为 {type_name}: {e}"
            )
    return result
