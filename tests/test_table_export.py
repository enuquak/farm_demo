# -*- coding: utf-8 -*-
"""导表引擎集成测试。"""
import os
import pytest
from openpyxl import Workbook

from tools.table_export import export_table, export_all


def _create_items_excel(path):
    """创建 items.xlsx 测试文件。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["item_id", "name", "type", "max_stack"])
    ws.append(["int", "str", "str", "int"])
    ws.append(["# 物品ID", "# 物品名称", "# 物品类型", "# 最大堆叠数"])
    ws.append([1, "木材", "RESOURCE", 99])
    ws.append([2, "石头", "RESOURCE", 99])
    wb.save(path)


def _create_crops_excel(path):
    """创建 crops.xlsx 测试文件。"""
    wb = Workbook()
    ws = wb.active
    ws.append(["crop_id", "name", "grow_time"])
    ws.append(["int", "str", "int"])
    ws.append(["# 作物ID", "# 作物名称", "# 生长时间"])
    ws.append([1, "小麦", 60])
    wb.save(path)


class TestExportTable:
    """单表导出测试。"""

    def test_export_success(self, tmp_path):
        """正常导出生成 Python 和 C++ 文件。"""
        # 创建目录结构
        tables_dir = tmp_path / "tables"
        tables_dir.mkdir()
        client_data = tmp_path / "scripts" / "client" / "data"
        client_data.mkdir(parents=True)
        server_data = tmp_path / "scripts" / "server" / "data"
        server_data.mkdir(parents=True)

        # 创建 Excel
        excel_path = str(tables_dir / "items.xlsx")
        _create_items_excel(excel_path)

        # 导出
        result = export_table(excel_path, tmp_path)

        assert result["success"] is True
        assert os.path.exists(result["python_path"])
        assert os.path.exists(result["cpp_path"])

        # 验证 Python 内容
        with open(result["python_path"], "r", encoding="utf-8") as f:
            content = f.read()
        assert "ITEM_DEFS" in content

    def test_export_validation_error(self, tmp_path):
        """校验失败时返回错误信息，不生成文件。"""
        tables_dir = tmp_path / "tables"
        tables_dir.mkdir()
        (tmp_path / "scripts" / "client" / "data").mkdir(parents=True)
        (tmp_path / "scripts" / "server" / "data").mkdir(parents=True)

        # 创建有错误的 Excel（重复主键）
        wb = Workbook()
        ws = wb.active
        ws.append(["item_id", "name"])
        ws.append(["int", "str"])
        ws.append(["# ID", "# 名称"])
        ws.append([1, "木材"])
        ws.append([1, "石头"])  # 重复
        excel_path = str(tables_dir / "bad.xlsx")
        wb.save(excel_path)

        result = export_table(excel_path, tmp_path)

        assert result["success"] is False
        assert len(result["errors"]) > 0


class TestExportAll:
    """全表导出测试。"""

    def test_export_all_success(self, tmp_path):
        """导出目录下所有 Excel 文件。"""
        tables_dir = tmp_path / "tables"
        tables_dir.mkdir()
        (tmp_path / "scripts" / "client" / "data").mkdir(parents=True)
        (tmp_path / "scripts" / "server" / "data").mkdir(parents=True)

        _create_items_excel(str(tables_dir / "items.xlsx"))
        _create_crops_excel(str(tables_dir / "crops.xlsx"))

        result = export_all(str(tables_dir), tmp_path)

        assert result["total"] == 2
        assert result["success_count"] == 2
        assert result["fail_count"] == 0

    def test_export_all_mixed(self, tmp_path):
        """部分成功部分失败。"""
        tables_dir = tmp_path / "tables"
        tables_dir.mkdir()
        (tmp_path / "scripts" / "client" / "data").mkdir(parents=True)
        (tmp_path / "scripts" / "server" / "data").mkdir(parents=True)

        # 好的表
        _create_items_excel(str(tables_dir / "items.xlsx"))

        # 坏的表（重复主键）
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "name"])
        ws.append(["int", "str"])
        ws.append(["# ID", "# 名称"])
        ws.append([1, "a"])
        ws.append([1, "b"])
        wb.save(str(tables_dir / "bad.xlsx"))

        result = export_all(str(tables_dir), tmp_path)

        assert result["total"] == 2
        assert result["success_count"] == 1
        assert result["fail_count"] == 1
